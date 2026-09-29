# handles loading whatever file the user gives us - could be a plain txt, a pdf (digital or scanned), or an image. uses PyMuPDF (fitz) for
# pdfs since it can read text AND turn pages into images without needing poppler installed separately

import os
import fitz  # PyMuPDF
import cv2
import numpy as np

SUPPORTED_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".tiff"}
SUPPORTED_PDF_EXT = {".pdf"}
SUPPORTED_TEXT_EXT = {".txt"}
SUPPORTED_DOCX_EXT = {".docx"}


def load_text_file(path: str) -> str:
# reads and returns the contents of a plain text file
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def load_docx_file(path: str) -> str:
# reads a .docx file and returns its visible text - paragraphs plus
# any text sitting inside tables, since a lot of resume templates
# (including two-column ones) put all their content in a table
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def load_image_file(path: str) -> np.ndarray:
# read image using opencv, fallback to PIL if opencv can't handle
    img = cv2.imread(path)
    if img is None:
        from PIL import Image
        pil_img = Image.open(path).convert("RGB")
        img = cv2.cvtColor(np.array(pil_img), cv2.COLOR_RGB2BGR)
    return img


def pdf_page_to_image(page: "fitz.Page", zoom: float = 2.0) -> np.ndarray:
# render a PDF page as an image
# a higher zoom gives more pixels, which usually improves OCR accuracy
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    if pix.n == 4:
        img = cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
    elif pix.n == 3:
        img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
    else:
        img = cv2.cvtColor(img, cv2.COLOR_GRAY2BGR)
    return img


def load_pdf(path: str) -> dict:
# it opens a pdf and grabs 2 things the text layes and rendered page images

    doc = fitz.open(path)
    digital_text_parts = []
    page_images = []
    for page in doc:
        digital_text_parts.append(page.get_text())
        page_images.append(pdf_page_to_image(page))
    doc.close()
    return {
        "digital_text": "\n".join(digital_text_parts).strip(),
        "page_images": page_images,
    }


def load_document(path: str) -> dict:
    # figures out file type from extension and loads it accordingly
    #  always returns same shape of dict so the rest of the code doesnt need to take care of the format 
    ext = os.path.splitext(path)[1].lower()
    if ext in SUPPORTED_TEXT_EXT:
        return {"type": "text", "digital_text": load_text_file(path), "page_images": []}
    if ext in SUPPORTED_PDF_EXT:
        result = load_pdf(path)
        return {"type": "pdf", **result}
    if ext in SUPPORTED_IMAGE_EXT:
        img = load_image_file(path)
        return {"type": "image", "digital_text": "", "page_images": [img]}
    if ext in SUPPORTED_DOCX_EXT:
        return {"type": "text", "digital_text": load_docx_file(path), "page_images": []}
    raise ValueError(f"Unsupported file type: {ext}")


def list_resume_files(folder: str) -> list:
    # returns all supported resume files in a folder, sorted by name
    exts = SUPPORTED_IMAGE_EXT | SUPPORTED_PDF_EXT | SUPPORTED_TEXT_EXT | SUPPORTED_DOCX_EXT
    files = []
    for fname in sorted(os.listdir(folder)):
        if os.path.splitext(fname)[1].lower() in exts:
            files.append(os.path.join(folder, fname))
    return files